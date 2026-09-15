"""Importação e utilitários de imóveis vendidos (planilha Resumo)."""

from __future__ import annotations

from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from typing import Any, BinaryIO, Iterable

from django.db import transaction
from django.utils import timezone

from controladoria.models import ImovelVendido

SHEET_NAME = 'Resumo'

HEADER_TO_FIELD = {
    'empreendimento ajustado': 'empreendimento_ajustado',
    'cod empreendimento': 'cod_empreendimento',
    'cód empreendimento': 'cod_empreendimento',
    'data da venda': 'data_venda',
    'contrato ajustado': 'contrato_ajustado',
    'situação': 'situacao',
    'situacao': 'situacao',
    'imóvel': 'imovel',
    'imovel': 'imovel',
    'cliente': 'cliente',
    'valor tabela': 'valor_tabela',
    'desconto': 'desconto',
    'valor venda': 'valor_venda',
    'valor liquidado': 'valor_liquidado',
    'situação do contrato': 'situacao_contrato',
    'situacao do contrato': 'situacao_contrato',
    'data de rescisão imobiliaria': 'data_rescisao_imobiliaria',
    'data de rescisao imobiliaria': 'data_rescisao_imobiliaria',
    'imobiliaria': 'imobiliaria',
    'imobiliária': 'imobiliaria',
    'corretor': 'corretor',
    'regra comissão': 'regra_comissao',
    'regra comissao': 'regra_comissao',
    'comissão': 'comissao',
    'comissao': 'comissao',
    'regra valor': 'regra_valor',
    'competência': 'competencia',
    'competencia': 'competencia',
}

DATE_FIELDS = {'data_venda', 'data_rescisao_imobiliaria', 'competencia'}
DECIMAL_FIELDS = {
    'valor_tabela',
    'desconto',
    'valor_venda',
    'valor_liquidado',
    'comissao',
}
UPDATE_FIELDS = [
    'empreendimento_ajustado',
    'cod_empreendimento',
    'data_venda',
    'situacao',
    'imovel',
    'cliente',
    'valor_tabela',
    'desconto',
    'valor_venda',
    'valor_liquidado',
    'situacao_contrato',
    'data_rescisao_imobiliaria',
    'imobiliaria',
    'corretor',
    'regra_comissao',
    'comissao',
    'regra_valor',
    'competencia',
    'atualizado_em',
]


class ImovelVendidoImportError(Exception):
    """Erro de leitura/importação da planilha."""


@dataclass
class ImportResultado:
    criados: int = 0
    atualizados: int = 0
    ignorados: int = 0
    erros: list[str] = field(default_factory=list)


def _normalizar_cabecalho(valor: Any) -> str:
    texto = str(valor or '').strip().lower()
    return ' '.join(texto.split())


def _para_texto(valor: Any, max_len: int | None = None) -> str:
    if valor is None:
        return ''
    if isinstance(valor, datetime):
        texto = valor.strftime('%d/%m/%Y')
    elif isinstance(valor, date):
        texto = valor.strftime('%d/%m/%Y')
    else:
        texto = str(valor).strip()
    if max_len is not None:
        return texto[:max_len]
    return texto


def _para_data(valor: Any) -> date | None:
    if valor is None or valor == '':
        return None
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    texto = str(valor).strip()
    if not texto:
        return None
    for fmt in ('%d/%m/%Y', '%Y-%m-%d', '%d-%m-%Y', '%d/%m/%y'):
        try:
            return datetime.strptime(texto, fmt).date()
        except ValueError:
            continue
    raise ValueError(f'Data inválida: {texto}')


def _para_decimal(valor: Any) -> Decimal | None:
    if valor is None or valor == '':
        return None
    if isinstance(valor, Decimal):
        return valor
    if isinstance(valor, (int, float)):
        return Decimal(str(valor))
    texto = str(valor).strip()
    if not texto:
        return None
    if ',' in texto and '.' in texto:
        texto = texto.replace('.', '').replace(',', '.')
    elif ',' in texto:
        texto = texto.replace(',', '.')
    try:
        return Decimal(texto)
    except InvalidOperation as exc:
        raise ValueError(f'Valor numérico inválido: {valor}') from exc


def _mapear_cabecalhos(headers: Iterable[Any]) -> dict[int, str]:
    mapeamento: dict[int, str] = {}
    for idx, header in enumerate(headers):
        chave = _normalizar_cabecalho(header)
        if not chave:
            continue
        # Remove acentos simples para matching mais tolerante.
        chave_ascii = (
            chave.replace('á', 'a').replace('à', 'a').replace('ã', 'a').replace('â', 'a')
            .replace('é', 'e').replace('ê', 'e')
            .replace('í', 'i')
            .replace('ó', 'o').replace('ô', 'o').replace('õ', 'o')
            .replace('ú', 'u')
            .replace('ç', 'c')
        )
        field = HEADER_TO_FIELD.get(chave) or HEADER_TO_FIELD.get(chave_ascii)
        if field:
            mapeamento[idx] = field
    if 'contrato_ajustado' not in mapeamento.values():
        raise ImovelVendidoImportError(
            'A planilha não possui a coluna obrigatória "Contrato Ajustado" na aba Resumo.'
        )
    return mapeamento


def _linha_para_dados(row: tuple[Any, ...], colunas: dict[int, str]) -> dict[str, Any]:
    dados: dict[str, Any] = {}
    for idx, field in colunas.items():
        valor = row[idx] if idx < len(row) else None
        if field in DATE_FIELDS:
            dados[field] = _para_data(valor)
        elif field in DECIMAL_FIELDS:
            dados[field] = _para_decimal(valor)
        else:
            max_len = ImovelVendido._meta.get_field(field).max_length
            dados[field] = _para_texto(valor, max_len=max_len)
    return dados


def ler_linhas_planilha(arquivo: BinaryIO, sheet_name: str = SHEET_NAME) -> list[dict[str, Any]]:
    try:
        import openpyxl
    except ImportError as exc:
        raise ImovelVendidoImportError(
            'Biblioteca openpyxl não instalada. Execute: pip install openpyxl'
        ) from exc

    try:
        wb = openpyxl.load_workbook(arquivo, read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - mensagem amigável ao usuário
        raise ImovelVendidoImportError(f'Não foi possível abrir a planilha: {exc}') from exc

    try:
        nome = next((s for s in wb.sheetnames if s.lower() == sheet_name.lower()), None)
        if not nome:
            raise ImovelVendidoImportError(
                f'Aba "{sheet_name}" não encontrada. Abas disponíveis: {", ".join(wb.sheetnames)}'
            )
        ws = wb[nome]
        rows = ws.iter_rows(values_only=True)
        try:
            header_row = next(rows)
        except StopIteration as exc:
            raise ImovelVendidoImportError('A planilha está vazia.') from exc

        colunas = _mapear_cabecalhos(header_row)
        linhas: list[dict[str, Any]] = []
        for num, row in enumerate(rows, start=2):
            if row is None or all(v is None or str(v).strip() == '' for v in row):
                continue
            try:
                dados = _linha_para_dados(tuple(row), colunas)
            except ValueError as exc:
                raise ImovelVendidoImportError(f'Linha {num}: {exc}') from exc
            contrato = (dados.get('contrato_ajustado') or '').strip()
            if not contrato:
                continue
            dados['contrato_ajustado'] = contrato
            linhas.append(dados)
        return linhas
    finally:
        wb.close()


def importar_imoveis_vendidos(arquivo: BinaryIO, sheet_name: str = SHEET_NAME) -> ImportResultado:
    """Importa/atualiza imóveis vendidos a partir da aba Resumo."""
    linhas = ler_linhas_planilha(arquivo, sheet_name=sheet_name)
    resultado = ImportResultado()
    if not linhas:
        resultado.ignorados = 0
        return resultado

    agora = timezone.now()
    contratos = [linha['contrato_ajustado'] for linha in linhas]
    existentes = {
        obj.contrato_ajustado: obj
        for obj in ImovelVendido.objects.filter(contrato_ajustado__in=contratos)
    }

    criar: list[ImovelVendido] = []
    atualizar: list[ImovelVendido] = []

    for dados in linhas:
        contrato = dados['contrato_ajustado']
        existente = existentes.get(contrato)
        if existente is None:
            criar.append(ImovelVendido(**dados))
            continue

        mudou = False
        for campo in UPDATE_FIELDS:
            if campo == 'atualizado_em':
                continue
            novo = dados.get(campo)
            if getattr(existente, campo) != novo:
                setattr(existente, campo, novo)
                mudou = True
        if mudou:
            existente.atualizado_em = agora
            atualizar.append(existente)
        else:
            resultado.ignorados += 1

    with transaction.atomic():
        if criar:
            ImovelVendido.objects.bulk_create(criar, batch_size=500)
            resultado.criados = len(criar)
        if atualizar:
            ImovelVendido.objects.bulk_update(atualizar, UPDATE_FIELDS, batch_size=500)
            resultado.atualizados = len(atualizar)

    return resultado
